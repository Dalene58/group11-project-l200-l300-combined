from __future__ import annotations

import csv
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
from matplotlib import pyplot as plt
from docx import Document
from docx.shared import Inches

ROOT = Path(__file__).resolve().parent
PLOTS_DIR = ROOT / "plots"
LOGS_DIR = ROOT / "logs"
DOCX_PATH = ROOT / "submission_report.docx"


def read_csv_rows(csv_name: str):
    path = LOGS_DIR / csv_name
    with path.open("r", newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader)


def build_metrics():
    return {
        "Runtime": ("runtime_vs_input.csv", "time_ns", "Runtime (ns)"),
        "Memory": ("memory_vs_input.csv", "memory_kb", "Memory (KB)"),
        "Completion": ("completion_vs_input.csv", "completed_requests", "Completed Requests"),
    }


def plot_metric(metric_name: str, csv_name: str, value_key: str, y_label: str):
    rows = read_csv_rows(csv_name)
    by_algorithm: dict[str, list[tuple[int, float]]] = {}

    for row in rows:
        algorithm = row["algorithm_name"].strip()
        input_size = int(float(row["input_size"]))
        value = float(row[value_key])
        by_algorithm.setdefault(algorithm, []).append((input_size, value))

    fig, ax = plt.subplots(figsize=(8, 5))
    for algorithm, points in by_algorithm.items():
        points.sort(key=lambda x: x[0])
        x = [p[0] for p in points]
        y = [p[1] for p in points]
        ax.plot(x, y, marker="o", linewidth=2.5, label=algorithm)

    ax.set_title(f"{metric_name} Comparison Across Algorithms")
    ax.set_xlabel("Input Size")
    ax.set_ylabel(y_label)
    ax.grid(True, alpha=0.3)
    ax.legend()
    plt.tight_layout()
    fig.savefig(PLOTS_DIR / f"{metric_name.lower()}_comparison.png", dpi=200)
    plt.close(fig)


def add_docx_images(doc: Document):
    doc.add_heading("Ghana Smart Waste-Management Optimizer", level=0)
    doc.add_paragraph(
        "This document contains the essential benchmark comparisons generated directly from the project log files. "
        "Only the key performance metrics are included to keep the submission focused and readable."
    )

    for metric_name, (_, _, _) in build_metrics().items():
        comparison_path = PLOTS_DIR / f"{metric_name.lower()}_comparison.png"
        if comparison_path.exists():
            doc.add_heading(metric_name, level=1)
            doc.add_picture(str(comparison_path), width=Inches(6.5))
            doc.add_paragraph(f"Figure: {metric_name} comparison between the main algorithms.")


def main():
    PLOTS_DIR.mkdir(exist_ok=True)

    for metric_name, (csv_name, value_key, y_label) in build_metrics().items():
        plot_metric(metric_name, csv_name, value_key, y_label)

    doc = Document()
    add_docx_images(doc)
    doc.save(DOCX_PATH)

    print(f"Generated comparison plots in: {PLOTS_DIR}")
    print(f"Updated document: {DOCX_PATH}")


if __name__ == "__main__":
    main()
